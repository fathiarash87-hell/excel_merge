import pandas as pd
from pathlib import Path

data_folder = Path("data")
frames = []

for file in sorted(data_folder.glob("*.xlsx")):
    df = pd.read_excel(file)
    df["Source"] = file.stem
    frames.append(df)

merged = pd.concat(frames, ignore_index=True)

summary = merged.groupby("Source")["salary"].agg(["sum", "mean", "min", "max"]).reset_index()

with pd.ExcelWriter("report.xlsx") as writer:
    merged.to_excel(writer, sheet_name="Merged", index=False)
    summary.to_excel(writer, sheet_name="Summary", index=False)

print("فایل report.xlsx ساخته شد")
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import BarChart, Reference, PieChart, LineChart3D

wb = load_workbook("report.xlsx")

# فرمت‌بندی هر دو شیت
for ws in wb.worksheets:
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4F6F78")
        cell.alignment = Alignment(horizontal="center")

    for column in ws.columns:
        max_len = max(len(str(c.value)) if c.value is not None else 0 for c in column)
        ws.column_dimensions[column[0].column_letter].width = max_len + 4

    ws.freeze_panes = "A2"

# نمودار در شیت Summary
ws = wb["Summary"]
chart = PieChart()
chart.title = "جمع حقوق هر ماه"
data = Reference(ws, min_col=2, min_row=1, max_row=ws.max_row)
cats = Reference(ws, min_col=1, min_row=2, max_row=ws.max_row)
chart.add_data(data, titles_from_data=True)
chart.set_categories(cats)
ws.add_chart(chart, "D6")

wb.save("report.xlsx")
print("فرمت‌بندی و نمودار اضافه شد")
